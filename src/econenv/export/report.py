"""A short research document: text, tables, figures and what produced them.

This is deliberately not a manuscript system. It answers one question a
researcher asks at the end of a session — *put the table, the chart and the
versions into one file I can send someone* — and stops there. Anything more is
a job for a real authoring tool, and pretending otherwise would produce a
fragile imitation of one.

The reproducibility block is the part that earns its place: a table and a figure
without the versions that made them are the thing referees cannot check.

    report = econenv.report("Consumption function")
    report.add_text("Quarterly US data, 1959-2009.")
    report.add_table(comparison, caption="Table 1. Baseline estimates")
    report.add_figures(execution.figures, caption="Figure 1. Residuals")
    report.add_snapshot()
    report.write("paper/report.docx")
"""

from __future__ import annotations

import datetime as _datetime
import html as _html
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Union

import pandas as pd

from ..exceptions import EconEnvError
from .figures import save_figures
from .tables import DEFAULT_STARS, star_note
from .writers import latex_escape, tex_engine, write_docx

#: What a report can be written as.
REPORT_FORMATS = ("html", "md", "tex", "docx", "pdf")


@dataclass
class _Block:
    """One piece of a report."""

    kind: str  # text | table | figure | snapshot
    payload: Any = None
    caption: Optional[str] = None
    note: Optional[str] = None


@dataclass
class Report:
    """A document being assembled. Add blocks in the order they should appear."""

    title: str = "Results"
    author: Optional[str] = None
    date: Optional[str] = None
    blocks: List[_Block] = field(default_factory=list)

    # ------------------------------------------------------------------ #
    # building
    # ------------------------------------------------------------------ #
    def add_text(self, text: str) -> Report:
        """A paragraph. Blank lines separate paragraphs."""
        self.blocks.append(_Block("text", str(text)))
        return self

    def add_table(
        self,
        obj: Any,
        *,
        caption: Optional[str] = None,
        style: str = "journal",
        digits: int = 4,
        stars: Sequence = DEFAULT_STARS,
    ) -> Report:
        """A model result, comparison, DataFrame or matrix as a table."""
        from . import _has_estimates, _to_frame

        frame = _to_frame(obj, style=style, stars=stars, digits=digits)
        note = star_note(stars) if (style == "journal" and _has_estimates(obj)) else None
        self.blocks.append(_Block("table", frame, caption, note))
        return self

    def add_figure(self, figure: Any, *, caption: Optional[str] = None) -> Report:
        """One captured figure."""
        self.blocks.append(_Block("figure", [figure], caption))
        return self

    def add_figures(self, figures: Sequence[Any], *, caption: Optional[str] = None) -> Report:
        """Several figures under one caption."""
        self.blocks.append(_Block("figure", list(figures), caption))
        return self

    def add_snapshot(self, snapshot: Optional[Dict[str, Any]] = None) -> Report:
        """The reproducibility block: EconEnv, Python and every engine's version.

        Taken live when not supplied. It never raises for an engine that is not
        installed — an absent engine is a fact about the machine, and a report
        that refuses to be written because Stata is missing is useless.
        """
        if snapshot is None:
            from .. import services

            try:
                snapshot = services.snapshot()
            except Exception as exc:  # a broken engine must not lose the report
                snapshot = {"error": f"snapshot unavailable: {exc}"}
        self.blocks.append(_Block("snapshot", snapshot))
        return self

    # ------------------------------------------------------------------ #
    # writing
    # ------------------------------------------------------------------ #
    def write(self, path: Union[str, Path], *, fmt: Optional[str] = None) -> Path:
        """Write the report. The format follows the extension unless *fmt* says."""
        target = Path(path)
        chosen = (fmt or target.suffix.lstrip(".") or "html").lower()
        if chosen not in REPORT_FORMATS:
            raise EconEnvError(
                f"Cannot write a report as {chosen!r}. Available: {', '.join(REPORT_FORMATS)}."
            )
        target.parent.mkdir(parents=True, exist_ok=True)

        if chosen == "html":
            target.write_text(self.to_html(target.parent), encoding="utf-8")
        elif chosen == "md":
            target.write_text(self.to_markdown(target.parent), encoding="utf-8")
        elif chosen == "tex":
            target.write_text(self.to_latex(target.parent), encoding="utf-8")
        elif chosen == "docx":
            self._write_docx(target)
        elif chosen == "pdf":
            self._write_pdf(target)
        return target

    # ------------------------------------------------------------------ #
    # renderers
    # ------------------------------------------------------------------ #
    def _heading(self) -> List[str]:
        parts = [self.title]
        if self.author:
            parts.append(self.author)
        parts.append(self.date or _datetime.date.today().isoformat())
        return parts

    def _figure_paths(self, directory: Path) -> Dict[int, List[Path]]:
        """Write every figure beside the report, once, and remember where."""
        written: Dict[int, List[Path]] = {}
        for index, block in enumerate(self.blocks):
            if block.kind != "figure" or not block.payload:
                continue
            written[index] = save_figures(block.payload, directory, prefix=f"report{index:02d}")
        return written

    def to_html(self, directory: Optional[Path] = None) -> str:
        directory = Path(directory or ".")
        figures = self._figure_paths(directory)
        title, *rest = self._heading()

        out = [
            "<!doctype html>",
            '<meta charset="utf-8">',
            f"<title>{_html.escape(title)}</title>",
            "<style>",
            "body{max-width:52em;margin:3em auto;padding:0 1.5em;"
            "font:16px/1.6 Georgia,'Times New Roman',serif;color:#1b1b1b}",
            "h1{font-size:1.7em;margin-bottom:.15em}",
            ".meta{color:#666;font-size:.85em;margin-bottom:2.5em}",
            "table{border-collapse:collapse;margin:1.2em 0;font-size:.92em}",
            "th,td{padding:.35em .8em;text-align:right}",
            "th:first-child,td:first-child{text-align:left}",
            "thead th{border-bottom:1px solid #333}",
            "table{border-top:2px solid #333;border-bottom:2px solid #333}",
            ".caption{font-weight:600;margin-top:1.6em}",
            ".note{font-size:.82em;color:#555;margin-top:-.6em}",
            "figure{margin:1.4em 0}img{max-width:100%}",
            "figcaption{font-size:.85em;color:#555}",
            ".snap{background:#f6f5f2;border:1px solid #e3e1db;padding:1em;"
            "font:13px/1.5 ui-monospace,Consolas,monospace;white-space:pre-wrap}",
            "</style>",
            f"<h1>{_html.escape(title)}</h1>",
            f'<p class="meta">{" &middot; ".join(_html.escape(r) for r in rest)}</p>',
        ]

        for index, block in enumerate(self.blocks):
            if block.kind == "text":
                for paragraph in str(block.payload).split("\n\n"):
                    out.append(f"<p>{_html.escape(paragraph.strip())}</p>")
            elif block.kind == "table":
                if block.caption:
                    out.append(f'<p class="caption">{_html.escape(block.caption)}</p>')
                out.append(block.payload.to_html(border=0, escape=True))
                if block.note:
                    out.append(f'<p class="note">{_html.escape(block.note)}</p>')
            elif block.kind == "figure":
                if block.caption:
                    out.append(f'<p class="caption">{_html.escape(block.caption)}</p>')
                for path in figures.get(index, []):
                    out.append(f'<figure><img src="{path.name}" alt="{path.stem}"></figure>')
            elif block.kind == "snapshot":
                out.append('<p class="caption">Reproducibility</p>')
                out.append(f'<div class="snap">{_html.escape(_snapshot_text(block.payload))}</div>')
        return "\n".join(out) + "\n"

    def to_markdown(self, directory: Optional[Path] = None) -> str:
        directory = Path(directory or ".")
        figures = self._figure_paths(directory)
        title, *rest = self._heading()

        out = [f"# {title}", "", " · ".join(rest), ""]
        for index, block in enumerate(self.blocks):
            if block.kind == "text":
                out += [str(block.payload), ""]
            elif block.kind == "table":
                if block.caption:
                    out += [f"**{block.caption}**", ""]
                out += [block.payload.to_markdown(), ""]
                if block.note:
                    out += [f"*{block.note}*", ""]
            elif block.kind == "figure":
                if block.caption:
                    out += [f"**{block.caption}**", ""]
                for path in figures.get(index, []):
                    out.append(f"![{path.stem}]({path.name})")
                out.append("")
            elif block.kind == "snapshot":
                out += ["## Reproducibility", "", "```", _snapshot_text(block.payload), "```", ""]
        return "\n".join(out)

    def to_latex(self, directory: Optional[Path] = None) -> str:
        directory = Path(directory or ".")
        figures = self._figure_paths(directory)
        title, *rest = self._heading()

        out = [
            r"\documentclass[11pt]{article}",
            r"\usepackage[margin=1in]{geometry}",
            r"\usepackage{booktabs}",
            r"\usepackage{graphicx}",
            r"\usepackage[T1]{fontenc}",
            r"\title{" + latex_escape(title) + "}",
        ]
        if self.author:
            out.append(r"\author{" + latex_escape(self.author) + "}")
        out += [
            r"\date{" + latex_escape(rest[-1]) + "}",
            r"\begin{document}",
            r"\maketitle",
        ]
        for index, block in enumerate(self.blocks):
            if block.kind == "text":
                out += [latex_escape(str(block.payload)), ""]
            elif block.kind == "table":
                out.append(r"\begin{table}[htbp]\centering")
                if block.caption:
                    out.append(r"\caption{" + latex_escape(block.caption) + "}")
                out.append(_latex_tabular(block.payload))
                if block.note:
                    out.append(
                        r"\begin{minipage}{\textwidth}\footnotesize "
                        + latex_escape(block.note)
                        + r"\end{minipage}"
                    )
                out.append(r"\end{table}")
            elif block.kind == "figure":
                for path in figures.get(index, []):
                    out.append(r"\begin{figure}[htbp]\centering")
                    out.append(r"\includegraphics[width=\textwidth]{" + path.name + "}")
                    if block.caption:
                        out.append(r"\caption{" + latex_escape(block.caption) + "}")
                    out.append(r"\end{figure}")
            elif block.kind == "snapshot":
                out += [
                    r"\section*{Reproducibility}",
                    r"\begin{verbatim}",
                    _snapshot_text(block.payload),
                    r"\end{verbatim}",
                ]
        out.append(r"\end{document}")
        return "\n".join(out) + "\n"

    def _write_docx(self, target: Path) -> None:
        try:
            import docx  # noqa: F401
        except ImportError as exc:
            raise EconEnvError(
                'Writing a .docx report needs python-docx: pip install "econenv[export]"'
            ) from exc
        from docx import Document
        from docx.shared import Inches

        document = Document()
        title, *rest = self._heading()
        document.add_heading(title, level=0)
        document.add_paragraph(" · ".join(rest))

        figures = self._figure_paths(target.parent)
        for index, block in enumerate(self.blocks):
            if block.kind == "text":
                document.add_paragraph(str(block.payload))
            elif block.kind == "table":
                if block.caption:
                    document.add_heading(block.caption, level=2)
                # Reuse the table writer rather than a second implementation.
                scratch = target.with_name(f"{target.stem}-t{index}.docx")
                write_docx(block.payload, scratch, note=block.note)
                _append_docx_body(document, scratch)
                scratch.unlink(missing_ok=True)
            elif block.kind == "figure":
                if block.caption:
                    document.add_heading(block.caption, level=2)
                for path in figures.get(index, []):
                    if path.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif"):
                        document.add_picture(str(path), width=Inches(6))
                    else:
                        # Word cannot place an SVG or PDF; name the file instead
                        # of dropping it silently.
                        document.add_paragraph(f"[figure: {path.name}]")
            elif block.kind == "snapshot":
                document.add_heading("Reproducibility", level=2)
                document.add_paragraph(_snapshot_text(block.payload))
        document.save(str(target))

    def _write_pdf(self, target: Path) -> None:
        import subprocess
        import tempfile

        engine = tex_engine()
        if engine is None:
            raise EconEnvError(
                "A PDF report needs a TeX engine (pdflatex, xelatex, lualatex or "
                "tectonic) on PATH, and there is none — it cannot be installed "
                "with pip.\n"
                "  Windows : MiKTeX (https://miktex.org) or TeX Live\n"
                "  macOS   : brew install --cask mactex-no-gui\n"
                "  Linux   : apt install texlive-latex-recommended\n"
                "Or write .docx, or .tex to compile with your own preamble."
            )
        with tempfile.TemporaryDirectory() as workdir:
            work = Path(workdir)
            source = work / "report.tex"
            source.write_text(self.to_latex(work), encoding="utf-8")
            command = (
                [engine, str(source)]
                if engine == "tectonic"
                else [engine, "-interaction=nonstopmode", "-halt-on-error", source.name]
            )
            completed = subprocess.run(
                command, cwd=work, capture_output=True, text=True, timeout=300
            )
            produced = work / "report.pdf"
            if not produced.exists():
                log = (completed.stdout or "") + (completed.stderr or "")
                errors = [ln for ln in log.splitlines() if ln.startswith("!")][:3]
                detail = "\n  ".join(errors) if errors else log.strip()[-400:]
                raise EconEnvError(f"{engine} could not typeset the report:\n  {detail}")
            shutil.copyfile(produced, target)

    def __str__(self) -> str:
        counts: Dict[str, int] = {}
        for block in self.blocks:
            counts[block.kind] = counts.get(block.kind, 0) + 1
        summary = ", ".join(f"{n} {k}" for k, n in counts.items()) or "empty"
        return f"Report({self.title!r}: {summary})"

    def _repr_html_(self) -> str:
        return self.to_html(Path("."))


def report(title: str = "Results", *, author: Optional[str] = None, **kwargs: Any) -> Report:
    """Start a report. See :class:`Report`."""
    return Report(title=title, author=author, **kwargs)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _snapshot_text(snapshot: Any) -> str:
    """The reproducibility block as flat, readable lines."""
    if isinstance(snapshot, str):
        return snapshot
    if not isinstance(snapshot, dict):
        return str(snapshot)

    lines: List[str] = []

    def walk(mapping: Dict[str, Any], indent: int = 0) -> None:
        for key, value in mapping.items():
            pad = "  " * indent
            if isinstance(value, dict):
                lines.append(f"{pad}{key}:")
                walk(value, indent + 1)
            elif isinstance(value, (list, tuple)):
                lines.append(f"{pad}{key}: {', '.join(str(v) for v in value)}")
            else:
                lines.append(f"{pad}{key}: {value}")

    walk(snapshot)
    return "\n".join(lines)


def _latex_tabular(frame: pd.DataFrame) -> str:
    """The booktabs tabular for a frame, without the surrounding table float."""
    import tempfile

    from .writers import write_latex

    with tempfile.TemporaryDirectory() as workdir:
        path = Path(workdir) / "t.tex"
        write_latex(frame, path, fragment=True)
        return path.read_text(encoding="utf-8").strip()


def _append_docx_body(document: Any, source: Path) -> None:
    """Copy the body of a one-table .docx into *document*."""
    import copy

    from docx import Document

    other = Document(str(source))
    for element in other.element.body:
        document.element.body.append(copy.deepcopy(element))
